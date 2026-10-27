"""CV upload validation and text extraction (PDF + DOCX, OCR fallback)."""
import io
import re
import zipfile

import config
from services.errors import (
    FileTooLargeError,
    UnreadableDocumentError,
    UnsupportedFormatError,
)

URL_RE = re.compile(r"(https?://[^\s)>\]]+|www\.[^\s)>\]]+)", re.I)


def detect_type(filename, data):
    """Check the REAL file type from its bytes, not just the extension."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in config.ALLOWED_EXTENSIONS:
        raise UnsupportedFormatError(
            "Unsupported file type. Please upload a PDF or DOCX file.",
            "Allowed formats: PDF, DOCX. Export your CV from Word/Google Docs as one of these.",
        )
    if data[:5] == b"%PDF-":
        real = "pdf"
    elif data[:2] == b"PK" and zipfile.is_zipfile(io.BytesIO(data)):
        real = "docx"
    else:
        real = None
    if real is None or real != ext:
        raise UnsupportedFormatError(
            "The file content does not match its extension, or the file is corrupt.",
            "Re-export your CV as a PDF or DOCX and try again.",
        )
    return real


def validate_upload(file_storage):
    if file_storage is None or not file_storage.filename:
        raise UnsupportedFormatError(
            "No file selected.", "Choose a PDF or DOCX CV to upload."
        )
    data = file_storage.read()
    if len(data) == 0:
        raise UnreadableDocumentError(
            "The uploaded file is empty.", "Check the file opens on your computer, then upload again."
        )
    if len(data) > config.MAX_UPLOAD_BYTES:
        raise FileTooLargeError(
            f"File is larger than the {config.MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.",
            "Compress the file or remove large images.",
        )
    kind = detect_type(file_storage.filename, data)
    return data, kind


def _extract_pdf(data):
    text_parts, links = [], []
    page_count = 0
    has_tables = False
    has_images = False
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                txt = page.extract_text() or ""
                text_parts.append(txt)
                try:
                    if page.find_tables():
                        has_tables = True
                except Exception:
                    pass
                if page.images:
                    has_images = True
                for h in page.hyperlinks or []:
                    uri = h.get("uri")
                    if uri:
                        links.append(uri)
    except Exception as exc:  # encrypted, corrupt, etc.
        msg = str(exc).lower()
        if "password" in msg or "encrypt" in msg:
            raise UnreadableDocumentError(
                "This PDF is password-protected.",
                "Remove the password and upload again.",
            )
        raise UnreadableDocumentError(
            "We couldn't read this PDF - it may be corrupt.",
            "Try re-exporting it, or upload a DOCX version.",
        )
    return "\n".join(text_parts), links, page_count, has_tables, has_images


def _ocr_pdf(data):
    """Optional OCR fallback. Returns '' if OCR libraries are not installed."""
    try:
        import pytesseract
        from pdf2image import convert_from_bytes

        pages = convert_from_bytes(data, dpi=200)
        return "\n".join(pytesseract.image_to_string(p) for p in pages)
    except Exception:
        return ""


def _extract_docx(data):
    try:
        import docx

        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise UnreadableDocumentError(
            "We couldn't read this DOCX file - it may be corrupt or password-protected.",
            "Open it in Word, save a fresh copy, and upload again.",
        )
    parts = [p.text for p in document.paragraphs]
    has_tables = len(document.tables) > 0
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text not in cells:
                    cells.append(cell.text)
            parts.append("  ".join(cells))
    links = []
    try:
        for rel in document.part.rels.values():
            if "hyperlink" in rel.reltype:
                links.append(rel._target)
    except Exception:
        pass
    has_images = len(document.inline_shapes) > 0
    pages = max(1, round(len(" ".join(parts).split()) / 450))
    return "\n".join(parts), links, pages, has_tables, has_images


def extract_text(data, kind):
    """Return a dict with text, links and layout hints. Raises friendly errors."""
    used_ocr = False
    if kind == "pdf":
        text, links, pages, has_tables, has_images = _extract_pdf(data)
        if len(text.strip()) < config.MIN_TEXT_CHARS:
            ocr = _ocr_pdf(data)
            if len(ocr.strip()) >= config.MIN_TEXT_CHARS:
                text, used_ocr = ocr, True
    else:
        text, links, pages, has_tables, has_images = _extract_docx(data)

    text = text.replace("\x00", "").strip()
    if len(text) < config.MIN_TEXT_CHARS:
        raise UnreadableDocumentError(
            "This document has no selectable text (it may be a scan or an image-only file).",
            "Try a DOCX, or a text-based PDF exported directly from Word/Google Docs.",
        )

    links = list(dict.fromkeys(links + URL_RE.findall(text)))
    return {
        "text": text,
        "links": links,
        "pages": pages,
        "has_tables": has_tables,
        "has_images": has_images,
        "used_ocr": used_ocr,
    }
